"""
Automated Test Suite for SIH 2026 Presentation Deck
Validates compliance with SIH rules and PRD guardrails.
"""

import os
import pytest
from pptx import Presentation

DECK_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "PRALAY_SIH_2026_Presentation.pptx"))

def test_deck_file_exists():
    assert os.path.exists(DECK_PATH), f"Deck not found at {DECK_PATH}"

def test_slide_count_is_exactly_six():
    prs = Presentation(DECK_PATH)
    assert len(prs.slides) == 6, f"SIH Rule Violation: Expected exactly 6 slides, found {len(prs.slides)}"

def test_widescreen_dimensions():
    prs = Presentation(DECK_PATH)
    # Check 16:9 ratio: 13.333 in x 7.5 in
    width_in = prs.slide_width.inches
    height_in = prs.slide_height.inches
    assert round(width_in, 2) == 13.33
    assert round(height_in, 2) == 7.50

def test_guardrail_keywords_present():
    prs = Presentation(DECK_PATH)
    full_text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    full_text.append(paragraph.text)

    deck_corpus = " ".join(full_text)

    required_keywords = [
        "26192",
        "Compound Risk",
        "Estimated Lead Time",
        "35 Mins",
        "WHERE IS THE RISK",
        "HOW SEVERE",
        "HOW SOON",
        "WHAT NOW",
        "GREEN",
        "YELLOW",
        "ORANGE",
        "RED",
        "DETECT EARLIER. IDENTIFY LOCALLY. ACT IN TIME.",
        "Ministry of Home Affairs",
        "NDRF",
        "Chamoli"
    ]

    for kw in required_keywords:
        assert kw.lower() in deck_corpus.lower(), f"Missing required keyword/guardrail: '{kw}'"

def test_speaker_notes_present_on_all_slides():
    prs = Presentation(DECK_PATH)
    for i, slide in enumerate(prs.slides):
        notes_tf = slide.notes_slide.notes_text_frame
        assert notes_tf and len(notes_tf.text.strip()) > 30, f"Slide {i+1} is missing detailed speaker notes!"
