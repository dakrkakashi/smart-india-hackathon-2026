"""
PRALAYADARSHI - Smart India Hackathon 2026 Presentation Deck Builder
Strict 6-Slide Deck conforming to SIH Guidelines and PRD Architecture
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Add scripts directory to path to load content
sys.path.insert(0, os.path.dirname(__file__))
from ppt_content import SLIDES_CONTENT

# -------------------------------------------------------------
# DESIGN SYSTEM PALETTE
# -------------------------------------------------------------
COLOR_BG_DARK = RGBColor(11, 31, 58)        # #0B1F3A Dark Navy
COLOR_BG_CARD = RGBColor(16, 42, 76)        # #102A4C Card Background
COLOR_BG_CARD_LIGHT = RGBColor(23, 56, 98)  # #173862 Elevated Card
COLOR_BORDER = RGBColor(40, 78, 128)        # #284E80 Subtle Border

COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_SLATE_200 = RGBColor(226, 232, 240)
COLOR_SLATE_300 = RGBColor(203, 213, 225)
COLOR_SLATE_400 = RGBColor(148, 163, 184)
COLOR_CYAN_ACCENT = RGBColor(56, 189, 248)  # #38BDF8 Sky Accent
COLOR_BLUE_ACCENT = RGBColor(96, 165, 250)  # #60A5FA Blue Accent

# 4-Tier Severity Colors
COLOR_GREEN = RGBColor(34, 197, 94)         # #22C55E
COLOR_YELLOW = RGBColor(250, 204, 21)       # #FACC15
COLOR_ORANGE = RGBColor(249, 115, 22)       # #F97316
COLOR_RED = RGBColor(239, 68, 68)           # #EF4444

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"


# -------------------------------------------------------------
# HELPER UTILITIES
# -------------------------------------------------------------
def set_slide_background(slide):
    """Sets full-bleed dark navy background."""
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5)
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = COLOR_BG_DARK
    bg_shape.line.fill.background()
    return bg_shape

def add_header(slide, title_text, subtitle_text, category_tag="SIH 2026 • PS 26192 • Disaster Management"):
    """Standardized top banner for slides 2-6."""
    tb_tag = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.3))
    tf_tag = tb_tag.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_right = tf_tag.margin_top = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = category_tag.upper()
    p_tag.font.name = FONT_HEADING
    p_tag.font.size = Pt(9.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_CYAN_ACCENT

    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.733), Inches(0.95))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
    
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE
    p_title.space_after = Pt(2)

    p_sub = tf_title.add_paragraph()
    p_sub.text = subtitle_text
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(11.5)
    p_sub.font.color.rgb = COLOR_SLATE_300

def add_card(slide, left, top, width, height, fill_color=COLOR_BG_CARD, border_color=COLOR_BORDER):
    """Creates a rounded rectangle card container."""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = fill_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
    else:
        card.line.fill.background()
    return card

def add_speaker_notes(slide, notes_text):
    """Adds speaker notes to the slide."""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text


# -------------------------------------------------------------
# SLIDE 1: TITLE PAGE
# -------------------------------------------------------------
def build_slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    c = SLIDES_CONTENT["slide_1"]

    # Top Pill Badge
    pill = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.75), Inches(6.8), Inches(0.4)
    )
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(18, 50, 92)
    pill.line.color.rgb = COLOR_CYAN_ACCENT
    pill.line.width = Pt(1)
    tf_p = pill.text_frame
    tf_p.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_pill = tf_p.paragraphs[0]
    p_pill.text = "SMART INDIA HACKATHON 2026  •  PROBLEM STATEMENT 26192"
    p_pill.alignment = PP_ALIGN.CENTER
    p_pill.font.name = FONT_HEADING
    p_pill.font.size = Pt(9.5)
    p_pill.font.bold = True
    p_pill.font.color.rgb = COLOR_CYAN_ACCENT

    # Main Project Title
    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(7.5), Inches(1.3))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = 0
    p_hero = tf_t.paragraphs[0]
    p_hero.text = c["title"]
    p_hero.font.name = FONT_HEADING
    p_hero.font.size = Pt(44)
    p_hero.font.bold = True
    p_hero.font.color.rgb = COLOR_WHITE
    p_hero.space_after = Pt(2)

    # Core Tagline
    p_tagline = tf_t.add_paragraph()
    p_tagline.text = c["tagline"]
    p_tagline.font.name = FONT_HEADING
    p_tagline.font.size = Pt(20)
    p_tagline.font.bold = True
    p_tagline.font.color.rgb = COLOR_CYAN_ACCENT

    # Subtitle Paragraph
    tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(2.75), Inches(7.2), Inches(0.8))
    tf_s = tb_sub.text_frame
    tf_s.word_wrap = True
    tf_s.margin_left = tf_s.margin_top = 0
    p_sub = tf_s.paragraphs[0]
    p_sub.text = c["subtitle"]
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = COLOR_SLATE_200
    p_sub.space_after = Pt(4)

    # 4-Tier Risk Scale Strip
    add_card(slide, Inches(0.8), Inches(3.7), Inches(7.2), Inches(0.65), fill_color=RGBColor(14, 38, 70))
    tiers = [
        ("GREEN: Nominal", COLOR_GREEN),
        ("YELLOW: Watch", COLOR_YELLOW),
        ("ORANGE: Prepare", COLOR_ORANGE),
        ("RED: Evacuate", COLOR_RED)
    ]
    for i, (txt, col) in enumerate(tiers):
        t_card = add_card(slide, Inches(0.9 + i * 1.76), Inches(3.78), Inches(1.68), Inches(0.48), fill_color=RGBColor(20, 48, 86), border_color=col)
        tf_tc = t_card.text_frame
        tf_tc.vertical_anchor = MSO_ANCHOR.MIDDLE
        ptc = tf_tc.paragraphs[0]
        ptc.text = txt
        ptc.alignment = PP_ALIGN.CENTER
        ptc.font.name = FONT_HEADING
        ptc.font.size = Pt(9)
        ptc.font.bold = True
        ptc.font.color.rgb = col

    # Problem Statement / Ministry Meta Card
    add_card(slide, Inches(0.8), Inches(4.6), Inches(7.2), Inches(2.3), fill_color=COLOR_BG_CARD)
    tb_meta = slide.shapes.add_textbox(Inches(1.0), Inches(4.75), Inches(6.8), Inches(2.0))
    tf_meta = tb_meta.text_frame
    tf_meta.word_wrap = True
    
    p1 = tf_meta.paragraphs[0]
    p1.text = "ORGANIZATION & DOMAIN CONTEXT"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(10)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_CYAN_ACCENT
    p1.space_after = Pt(4)

    lines = [
        ("Organization: ", "Ministry of Home Affairs | NDRF / Disaster Management Division"),
        ("Problem Statement: ", "PS 26192 — Flash Flood Prediction System for Hilly Regions"),
        ("Theme / Category: ", "Disaster Management  •  Software Category"),
        ("Pilot Basin Focus: ", "Chamoli & Rudraprayag Districts, Uttarakhand (Himalayan Catchments)")
    ]
    for lbl, val in lines:
        p = tf_meta.add_paragraph()
        run_l = p.add_run()
        run_l.text = lbl
        run_l.font.bold = True
        run_l.font.size = Pt(10.5)
        run_l.font.color.rgb = COLOR_WHITE
        run_v = p.add_run()
        run_v.text = val
        run_v.font.size = Pt(10.5)
        run_v.font.color.rgb = COLOR_SLATE_300
        p.space_after = Pt(3)

    # Right Hero Callout Card: "35 MINUTES"
    add_card(slide, Inches(8.35), Inches(1.3), Inches(4.18), Inches(5.6), fill_color=RGBColor(16, 44, 82), border_color=COLOR_RED)
    tb_rc = slide.shapes.add_textbox(Inches(8.6), Inches(1.55), Inches(3.68), Inches(5.1))
    tf_rc = tb_rc.text_frame
    tf_rc.word_wrap = True
    
    prc_tag = tf_rc.paragraphs[0]
    prc_tag.text = "THE NORTH STAR LEAD TIME"
    prc_tag.font.name = FONT_HEADING
    prc_tag.font.size = Pt(11)
    prc_tag.font.bold = True
    prc_tag.font.color.rgb = COLOR_RED
    prc_tag.alignment = PP_ALIGN.CENTER
    prc_tag.space_after = Pt(6)

    prc_stat = tf_rc.add_paragraph()
    prc_stat.text = "35 Mins"
    prc_stat.font.name = FONT_HEADING
    prc_stat.font.size = Pt(48)
    prc_stat.font.bold = True
    prc_stat.font.color.rgb = COLOR_WHITE
    prc_stat.alignment = PP_ALIGN.CENTER
    prc_stat.space_after = Pt(4)

    prc_label = tf_rc.add_paragraph()
    prc_label.text = "Estimated Lead Time Window"
    prc_label.font.name = FONT_HEADING
    prc_label.font.size = Pt(13)
    prc_label.font.bold = True
    prc_label.font.color.rgb = COLOR_YELLOW
    prc_label.alignment = PP_ALIGN.CENTER
    prc_label.space_after = Pt(12)

    bullets = [
        "Transforms chaotic panic into structured, orderly village evacuation.",
        "Dynamic physics calculation based on river surge rate of rise & rainfall spikes.",
        "Directly triggers plain-language Citizen Card dispatch to local feature phones.",
        "Routes families to high-ground shelters strictly outside gorge flood plains."
    ]
    for b in bullets:
        pb = tf_rc.add_paragraph()
        pb.text = "• " + b
        pb.font.name = FONT_BODY
        pb.font.size = Pt(10.5)
        pb.font.color.rgb = COLOR_SLATE_200
        pb.space_after = Pt(6)

    add_speaker_notes(slide, c["speaker_notes"])


# -------------------------------------------------------------
# SLIDE 2: PROPOSED SOLUTION & NORTH STAR
# -------------------------------------------------------------
def build_slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    c = SLIDES_CONTENT["slide_2"]

    add_header(slide, c["title"], c["subtitle"])

    col_w = Inches(5.72)
    h_col = Inches(2.3)
    
    # Left Column: Problem
    add_card(slide, Inches(0.8), Inches(1.8), col_w, h_col, fill_color=COLOR_BG_CARD, border_color=RGBColor(239, 68, 68))
    tb_p = slide.shapes.add_textbox(Inches(0.95), Inches(1.9), col_w - Inches(0.3), h_col - Inches(0.2))
    tf_p = tb_p.text_frame
    tf_p.word_wrap = True
    pp_h = tf_p.paragraphs[0]
    pp_h.text = "CURRENT SYSTEM SHORTCOMINGS (THE PROBLEM)"
    pp_h.font.name = FONT_HEADING
    pp_h.font.size = Pt(11)
    pp_h.font.bold = True
    pp_h.font.color.rgb = COLOR_RED
    pp_h.space_after = Pt(4)

    for pt in c["columns"][0]["points"]:
        p = tf_p.add_paragraph()
        p.text = "✕ " + pt
        p.font.name = FONT_BODY
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_SLATE_300
        p.space_after = Pt(2.5)

    # Right Column: Solution
    add_card(slide, Inches(6.8), Inches(1.8), col_w, h_col, fill_color=COLOR_BG_CARD, border_color=COLOR_GREEN)
    tb_s = slide.shapes.add_textbox(Inches(6.95), Inches(1.9), col_w - Inches(0.3), h_col - Inches(0.2))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True
    ps_h = tf_s.paragraphs[0]
    ps_h.text = "THE PRALAY INNOVATION (OUR SOLUTION)"
    ps_h.font.name = FONT_HEADING
    ps_h.font.size = Pt(11)
    ps_h.font.bold = True
    ps_h.font.color.rgb = COLOR_GREEN
    ps_h.space_after = Pt(4)

    for pt in c["columns"][1]["points"]:
        p = tf_s.add_paragraph()
        p.text = "✓ " + pt
        p.font.name = FONT_BODY
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_WHITE
        p.space_after = Pt(2.5)

    # 4 North Star Cards
    card_w = Inches(2.78)
    card_h = Inches(1.7)
    colors = {
        "BLUE": COLOR_CYAN_ACCENT,
        "ORANGE": COLOR_ORANGE,
        "RED": COLOR_RED,
        "GREEN": COLOR_GREEN
    }

    for i, nsc in enumerate(c["north_star_cards"]):
        left = Inches(0.8 + i * 2.98)
        col_accent = colors.get(nsc["color"], COLOR_CYAN_ACCENT)
        add_card(slide, left, Inches(4.25), card_w, card_h, fill_color=COLOR_BG_CARD_LIGHT, border_color=col_accent)
        tb_ns = slide.shapes.add_textbox(left + Inches(0.12), Inches(4.35), card_w - Inches(0.24), card_h - Inches(0.2))
        tf_ns = tb_ns.text_frame
        tf_ns.word_wrap = True
        
        p_q = tf_ns.paragraphs[0]
        p_q.text = nsc["q"]
        p_q.font.name = FONT_HEADING
        p_q.font.size = Pt(10.5)
        p_q.font.bold = True
        p_q.font.color.rgb = col_accent
        p_q.space_after = Pt(4)

        p_a = tf_ns.add_paragraph()
        p_a.text = nsc["ans"]
        p_a.font.name = FONT_BODY
        p_a.font.size = Pt(9.5)
        p_a.font.color.rgb = COLOR_SLATE_200

    # Bottom Dual-Interface Strip
    add_card(slide, Inches(0.8), Inches(6.1), Inches(11.733), Inches(0.95), fill_color=RGBColor(14, 38, 70), border_color=COLOR_BORDER)
    tb_strip = slide.shapes.add_textbox(Inches(0.95), Inches(6.15), Inches(11.433), Inches(0.85))
    tf_strip = tb_strip.text_frame
    tf_strip.word_wrap = True
    
    p_auth = tf_strip.paragraphs[0]
    r1 = p_auth.add_run()
    r1.text = "DUAL-INTERFACE PARADIGM  •  "
    r1.font.bold = True
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = COLOR_CYAN_ACCENT
    r2 = p_auth.add_run()
    r2.text = c["dual_interface"]["authority"]
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = COLOR_SLATE_300

    p_cit = tf_strip.add_paragraph()
    r3 = p_cit.add_run()
    r3.text = "CITIZEN ALERT CARD (PRD §18)  •  "
    r3.font.bold = True
    r3.font.size = Pt(9.5)
    r3.font.color.rgb = COLOR_GREEN
    r4 = p_cit.add_run()
    r4.text = c["dual_interface"]["citizen"]
    r4.font.size = Pt(9.5)
    r4.font.color.rgb = COLOR_WHITE

    add_speaker_notes(slide, c["speaker_notes"])


# -------------------------------------------------------------
# SLIDE 3: TECHNICAL APPROACH & PIPELINE
# -------------------------------------------------------------
def build_slide_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    c = SLIDES_CONTENT["slide_3"]

    add_header(slide, c["title"], c["subtitle"])

    # 5-Step Pipeline Cards
    pipe_w = Inches(2.22)
    pipe_h = Inches(1.4)
    for i, step in enumerate(c["pipeline_steps"]):
        left = Inches(0.8 + i * 2.38)
        add_card(slide, left, Inches(1.8), pipe_w, pipe_h, fill_color=COLOR_BG_CARD_LIGHT, border_color=COLOR_CYAN_ACCENT)
        tb_st = slide.shapes.add_textbox(left + Inches(0.1), Inches(1.88), pipe_w - Inches(0.2), pipe_h - Inches(0.16))
        tf_st = tb_st.text_frame
        tf_st.word_wrap = True
        
        p1 = tf_st.paragraphs[0]
        p1.text = f"STEP {step['num']}: {step['name']}"
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_CYAN_ACCENT
        p1.space_after = Pt(2)

        p2 = tf_st.add_paragraph()
        p2.text = step["desc"]
        p2.font.name = FONT_BODY
        p2.font.size = Pt(9)
        p2.font.color.rgb = COLOR_SLATE_200

    # Middle Row Left: Compound Risk Formula Card
    add_card(slide, Inches(0.8), Inches(3.38), Inches(6.8), Inches(2.2), fill_color=COLOR_BG_CARD, border_color=COLOR_BLUE_ACCENT)
    tb_form = slide.shapes.add_textbox(Inches(1.0), Inches(3.5), Inches(6.4), Inches(1.95))
    tf_form = tb_form.text_frame
    tf_form.word_wrap = True
    
    pf1 = tf_form.paragraphs[0]
    pf1.text = "DETERMINISTIC COMPOUND RISK FORMULATION (PRD §9 & §10)"
    pf1.font.name = FONT_HEADING
    pf1.font.size = Pt(10.5)
    pf1.font.bold = True
    pf1.font.color.rgb = COLOR_CYAN_ACCENT
    pf1.space_after = Pt(6)

    pf2 = tf_form.add_paragraph()
    pf2.text = c["compound_formula"]["formula"]
    pf2.font.name = "Consolas"
    pf2.font.size = Pt(10.5)
    pf2.font.bold = True
    pf2.font.color.rgb = COLOR_YELLOW
    pf2.space_after = Pt(6)

    pf3 = tf_form.add_paragraph()
    pf3.text = c["compound_formula"]["rationale"]
    pf3.font.name = FONT_BODY
    pf3.font.size = Pt(9.5)
    pf3.font.color.rgb = COLOR_SLATE_300
    pf3.space_after = Pt(4)

    pf4 = tf_form.add_paragraph()
    pf4.text = "• Flood Factor (0-100): River level + rate of rise + rainfall accumulation\n• Landslide Factor (0-100): Saturation pore pressure + slope steepness + trigger rain"
    pf4.font.name = FONT_BODY
    pf4.font.size = Pt(9)
    pf4.font.color.rgb = COLOR_SLATE_400

    # Middle Row Right: Engineering Guardrails Card
    add_card(slide, Inches(7.8), Inches(3.38), Inches(4.733), Inches(2.2), fill_color=COLOR_BG_CARD, border_color=COLOR_ORANGE)
    tb_g = slide.shapes.add_textbox(Inches(8.0), Inches(3.5), Inches(4.333), Inches(1.95))
    tf_g = tb_g.text_frame
    tf_g.word_wrap = True
    
    pg1 = tf_g.paragraphs[0]
    pg1.text = "ENGINEERING SAFETY & ACCURACY GUARDRAILS"
    pg1.font.name = FONT_HEADING
    pg1.font.size = Pt(10.5)
    pg1.font.bold = True
    pg1.font.color.rgb = COLOR_ORANGE
    pg1.space_after = Pt(6)

    for g in c["validation_guardrails"]:
        p = tf_g.add_paragraph()
        p.text = "🛡 " + g
        p.font.name = FONT_BODY
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_SLATE_200
        p.space_after = Pt(4)

    # Bottom Tech Stack Chips Bar
    add_card(slide, Inches(0.8), Inches(5.75), Inches(11.733), Inches(1.2), fill_color=RGBColor(14, 38, 70), border_color=COLOR_BORDER)
    tb_stk = slide.shapes.add_textbox(Inches(1.0), Inches(5.82), Inches(11.333), Inches(1.05))
    tf_stk = tb_stk.text_frame
    tf_stk.word_wrap = True
    
    pstk_lbl = tf_stk.paragraphs[0]
    pstk_lbl.text = "PRODUCTION ARCHITECTURE & CORE TECHNOLOGY STACK"
    pstk_lbl.font.name = FONT_HEADING
    pstk_lbl.font.size = Pt(9.5)
    pstk_lbl.font.bold = True
    pstk_lbl.font.color.rgb = COLOR_CYAN_ACCENT
    pstk_lbl.space_after = Pt(4)

    pstk_val = tf_stk.add_paragraph()
    chips_text = "   |   ".join(c["stack_chips"])
    pstk_val.text = chips_text
    pstk_val.font.name = FONT_HEADING
    pstk_val.font.size = Pt(10)
    pstk_val.font.bold = True
    pstk_val.font.color.rgb = COLOR_WHITE
    pstk_val.space_after = Pt(2)

    pstk_sub = tf_stk.add_paragraph()
    pstk_sub.text = "Async Non-blocking API Core • PostGIS Spatial Geometries • TimescaleDB Telemetry Hypertables • Edge IoT Telemetry"
    pstk_sub.font.name = FONT_BODY
    pstk_sub.font.size = Pt(8.5)
    pstk_sub.font.color.rgb = COLOR_SLATE_400

    add_speaker_notes(slide, c["speaker_notes"])


# -------------------------------------------------------------
# SLIDE 4: FEASIBILITY & ROADMAP
# -------------------------------------------------------------
def build_slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    c = SLIDES_CONTENT["slide_4"]

    add_header(slide, c["title"], c["subtitle"])

    card_w = Inches(3.78)
    card_h = Inches(2.2)
    for i, ph in enumerate(c["roadmap_phases"]):
        left = Inches(0.8 + i * 3.98)
        border_c = COLOR_GREEN if i == 0 else COLOR_BLUE_ACCENT
        add_card(slide, left, Inches(1.8), card_w, card_h, fill_color=COLOR_BG_CARD, border_color=border_c)
        tb_ph = slide.shapes.add_textbox(left + Inches(0.15), Inches(1.9), card_w - Inches(0.3), card_h - Inches(0.2))
        tf_ph = tb_ph.text_frame
        tf_ph.word_wrap = True
        
        p1 = tf_ph.paragraphs[0]
        p1.text = ph["phase"]
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_WHITE
        p1.space_after = Pt(2)

        p2 = tf_ph.add_paragraph()
        p2.text = ph["timeline"]
        p2.font.name = FONT_HEADING
        p2.font.size = Pt(9.5)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_CYAN_ACCENT if i != 0 else COLOR_GREEN
        p2.space_after = Pt(6)

        p3 = tf_ph.add_paragraph()
        p3.text = ph["details"]
        p3.font.name = FONT_BODY
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = COLOR_SLATE_200

    # Risk Mitigation Matrix
    add_card(slide, Inches(0.8), Inches(4.2), Inches(8.2), Inches(2.6), fill_color=COLOR_BG_CARD, border_color=COLOR_BORDER)
    tb_mit = slide.shapes.add_textbox(Inches(1.0), Inches(4.3), Inches(7.8), Inches(2.4))
    tf_mit = tb_mit.text_frame
    tf_mit.word_wrap = True
    
    pm_h = tf_mit.paragraphs[0]
    pm_h.text = "FIELD RISK ASSESSMENT & RESILIENCE MITIGATION MATRIX"
    pm_h.font.name = FONT_HEADING
    pm_h.font.size = Pt(10.5)
    pm_h.font.bold = True
    pm_h.font.color.rgb = COLOR_CYAN_ACCENT
    pm_h.space_after = Pt(4)

    for item in c["risk_mitigation"]:
        p = tf_mit.add_paragraph()
        r_risk = p.add_run()
        r_risk.text = "▲ " + item["risk"] + ": "
        r_risk.font.bold = True
        r_risk.font.size = Pt(9.5)
        r_risk.font.color.rgb = COLOR_YELLOW
        
        r_mit = p.add_run()
        r_mit.text = item["mitigation"]
        r_mit.font.size = Pt(9.5)
        r_mit.font.color.rgb = COLOR_SLATE_200
        p.space_after = Pt(2.5)

    # Cost & Viability Card
    add_card(slide, Inches(9.2), Inches(4.2), Inches(3.333), Inches(2.6), fill_color=COLOR_BG_CARD_LIGHT, border_color=COLOR_GREEN)
    tb_cost = slide.shapes.add_textbox(Inches(9.35), Inches(4.35), Inches(3.033), Inches(2.3))
    tf_cost = tb_cost.text_frame
    tf_cost.word_wrap = True
    
    pc_h = tf_cost.paragraphs[0]
    pc_h.text = "ECONOMIC VIABILITY"
    pc_h.font.name = FONT_HEADING
    pc_h.font.size = Pt(11)
    pc_h.font.bold = True
    pc_h.font.color.rgb = COLOR_GREEN
    pc_h.space_after = Pt(4)

    pc_stat = tf_cost.add_paragraph()
    pc_stat.text = "₹15,000"
    pc_stat.font.name = FONT_HEADING
    pc_stat.font.size = Pt(28)
    pc_stat.font.bold = True
    pc_stat.font.color.rgb = COLOR_WHITE
    pc_stat.space_after = Pt(2)

    pc_desc = tf_cost.add_paragraph()
    pc_desc.text = "Avg. CapEx per Solar IoT Monitoring Node"
    pc_desc.font.name = FONT_HEADING
    pc_desc.font.size = Pt(9.5)
    pc_desc.font.bold = True
    pc_desc.font.color.rgb = COLOR_YELLOW
    pc_desc.space_after = Pt(6)

    pc_bullets = [
        "10x more affordable than imported seismic/Doppler systems.",
        "Solar + LiFePO4 battery with 5-day cloudy backup.",
        "Uses open standards (MQTT / LoRaWAN / GSM 4G)."
    ]
    for b in pc_bullets:
        p = tf_cost.add_paragraph()
        p.text = "• " + b
        p.font.name = FONT_BODY
        p.font.size = Pt(8.5)
        p.font.color.rgb = COLOR_SLATE_200
        p.space_after = Pt(2)

    add_speaker_notes(slide, c["speaker_notes"])


# -------------------------------------------------------------
# SLIDE 5: IMPACT & BENEFITS
# -------------------------------------------------------------
def build_slide_5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    c = SLIDES_CONTENT["slide_5"]

    add_header(slide, c["title"], c["subtitle"])

    stat_w = Inches(2.78)
    stat_h = Inches(1.7)
    for i, m in enumerate(c["target_metrics"]):
        left = Inches(0.8 + i * 2.98)
        add_card(slide, left, Inches(1.8), stat_w, stat_h, fill_color=COLOR_BG_CARD_LIGHT, border_color=COLOR_CYAN_ACCENT)
        tb_st = slide.shapes.add_textbox(left + Inches(0.1), Inches(1.9), stat_w - Inches(0.2), stat_h - Inches(0.2))
        tf_st = tb_st.text_frame
        tf_st.word_wrap = True
        
        p_num = tf_st.paragraphs[0]
        p_num.text = m["stat"]
        p_num.font.name = FONT_HEADING
        p_num.font.size = Pt(24)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_WHITE
        p_num.alignment = PP_ALIGN.CENTER
        p_num.space_after = Pt(2)

        p_lbl = tf_st.add_paragraph()
        p_lbl.text = m["label"]
        p_lbl.font.name = FONT_HEADING
        p_lbl.font.size = Pt(10.5)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = COLOR_YELLOW
        p_lbl.alignment = PP_ALIGN.CENTER
        p_lbl.space_after = Pt(2)

        p_sub = tf_st.add_paragraph()
        p_sub.text = m["sub"]
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(8.5)
        p_sub.font.color.rgb = COLOR_SLATE_300
        p_sub.alignment = PP_ALIGN.CENTER

    ben_w = Inches(3.78)
    ben_h = Inches(2.45)
    for i, sb in enumerate(c["stakeholder_benefits"]):
        left = Inches(0.8 + i * 3.98)
        add_card(slide, left, Inches(3.68), ben_w, ben_h, fill_color=COLOR_BG_CARD, border_color=COLOR_BLUE_ACCENT)
        tb_ben = slide.shapes.add_textbox(left + Inches(0.15), Inches(3.8), ben_w - Inches(0.3), ben_h - Inches(0.25))
        tf_ben = tb_ben.text_frame
        tf_ben.word_wrap = True
        
        p_role = tf_ben.paragraphs[0]
        p_role.text = sb["role"].upper()
        p_role.font.name = FONT_HEADING
        p_role.font.size = Pt(11)
        p_role.font.bold = True
        p_role.font.color.rgb = COLOR_CYAN_ACCENT
        p_role.space_after = Pt(6)

        p_b = tf_ben.add_paragraph()
        p_b.text = sb["benefit"]
        p_b.font.name = FONT_BODY
        p_b.font.size = Pt(9.5)
        p_b.font.color.rgb = COLOR_SLATE_200

    add_card(slide, Inches(0.8), Inches(6.3), Inches(11.733), Inches(0.75), fill_color=RGBColor(18, 48, 88), border_color=COLOR_GREEN)
    tb_man = slide.shapes.add_textbox(Inches(0.95), Inches(6.35), Inches(11.433), Inches(0.65))
    tf_man = tb_man.text_frame
    tf_man.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_man = tf_man.paragraphs[0]
    p_man.text = c["closing_callout"]
    p_man.alignment = PP_ALIGN.CENTER
    p_man.font.name = FONT_HEADING
    p_man.font.size = Pt(15)
    p_man.font.bold = True
    p_man.font.color.rgb = COLOR_GREEN

    add_speaker_notes(slide, c["speaker_notes"])


# -------------------------------------------------------------
# SLIDE 6: RESEARCH & REFERENCES
# -------------------------------------------------------------
def build_slide_6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    c = SLIDES_CONTENT["slide_6"]

    add_header(slide, c["title"], c["subtitle"])

    qw = Inches(5.72)
    qh = Inches(2.25)
    quads = c["quadrants"]

    coords = [
        (Inches(0.8), Inches(1.8)),
        (Inches(6.8), Inches(1.8)),
        (Inches(0.8), Inches(4.2)),
        (Inches(6.8), Inches(4.2))
    ]

    for i, (left, top) in enumerate(coords):
        q = quads[i]
        add_card(slide, left, top, qw, qh, fill_color=COLOR_BG_CARD, border_color=COLOR_BORDER)
        tb_q = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.12), qw - Inches(0.36), qh - Inches(0.24))
        tf_q = tb_q.text_frame
        tf_q.word_wrap = True
        
        p_qh = tf_q.paragraphs[0]
        p_qh.text = f"{i+1}. {q['title'].upper()}"
        p_qh.font.name = FONT_HEADING
        p_qh.font.size = Pt(10.5)
        p_qh.font.bold = True
        p_qh.font.color.rgb = COLOR_CYAN_ACCENT
        p_qh.space_after = Pt(4)

        for item in q["items"]:
            pi = tf_q.add_paragraph()
            pi.text = "• " + item
            pi.font.name = FONT_BODY
            pi.font.size = Pt(9)
            pi.font.color.rgb = COLOR_SLATE_200
            pi.space_after = Pt(2)

    tb_foot = slide.shapes.add_textbox(Inches(0.8), Inches(6.65), Inches(11.733), Inches(0.4))
    tf_foot = tb_foot.text_frame
    tf_foot.word_wrap = True
    p_foot = tf_foot.paragraphs[0]
    p_foot.text = c["footer_note"]
    p_foot.alignment = PP_ALIGN.CENTER
    p_foot.font.name = FONT_HEADING
    p_foot.font.size = Pt(9.5)
    p_foot.font.color.rgb = COLOR_SLATE_400

    add_speaker_notes(slide, c["speaker_notes"])


# -------------------------------------------------------------
# MAIN BUILD & QA ORCHESTRATOR
# -------------------------------------------------------------
def build_presentation(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    print("Building Slide 1: Title Page...")
    build_slide_1(prs)

    print("Building Slide 2: Proposed Solution & North Star...")
    build_slide_2(prs)

    print("Building Slide 3: Technical Approach & Pipeline...")
    build_slide_3(prs)

    print("Building Slide 4: Feasibility & Roadmap...")
    build_slide_4(prs)

    print("Building Slide 5: Impact & Benefits...")
    build_slide_5(prs)

    print("Building Slide 6: Research & References...")
    build_slide_6(prs)

    slide_count = len(prs.slides)
    assert slide_count == 6, f"ERROR: Expected exactly 6 slides per SIH rules, got {slide_count}"

    prs.save(output_path)
    print(f"[SUCCESS] 6-Slide Presentation generated at: {output_path}")

    # Also save a copy in pralay/ directory
    alt_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "PRALAY_SIH_2026_Presentation.pptx"))
    prs.save(alt_path)
    print(f"[SUCCESS] Secondary copy generated at: {alt_path}")
    return True

if __name__ == "__main__":
    out_root = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "PRALAY_SIH_2026_Presentation.pptx"))
    build_presentation(out_root)
